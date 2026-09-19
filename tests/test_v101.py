from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.ats_checker import run_ats_file_check


def test_home_is_tetris_board_and_legacy_home_blocks_are_removed() -> None:
    index = (ROOT / "ui" / "index.html").read_text(encoding="utf-8")
    css = (ROOT / "ui" / "css" / "product_polish.css").read_text(encoding="utf-8")

    assert "cvm-v101-home-board" in index
    assert 'id="newCvBtn"' in index
    assert 'id="cvmLatexHomeBtn"' in index
    assert 'id="openProjectBtn"' in index
    assert 'id="demoBtn"' in index
    assert "create_cv_button.png" in css
    assert "customize_button.png" in css
    assert "clip-path: polygon" in css
    assert "CVM WORKFLOW" not in index
    assert 'id="cvmModifyCvBtn"' not in index


def test_recent_projects_moved_to_sidebar() -> None:
    index = (ROOT / "ui" / "index.html").read_text(encoding="utf-8")

    assert 'id="openRecentBtn"' in index
    assert 'id="page-recent-manager"' in index
    assert 'id="cvmRecentProjectCards"' in index
    assert 'class="cvm-v05-panel"' not in index


def test_dark_only_and_connection_removed() -> None:
    index = (ROOT / "ui" / "index.html").read_text(encoding="utf-8")
    v101 = (ROOT / "ui" / "js" / "v101_core.js").read_text(encoding="utf-8")

    assert 'data-theme="dark"' in index
    assert 'id="themeToggle"' not in index
    assert 'settingsThemeBtn' not in index
    assert 'data-page="settings"' not in index
    assert 'id="connectionBadge"' not in index
    assert 'return "dark"' in v101


def test_ats_requires_pdf_or_docx_selection() -> None:
    index = (ROOT / "ui" / "index.html").read_text(encoding="utf-8")
    bridge = (ROOT / "src" / "bridge.py").read_text(encoding="utf-8")
    v101 = (ROOT / "ui" / "js" / "v101_core.js").read_text(encoding="utf-8")

    assert 'id="cvmAtsChooseFileBtn"' in index
    assert 'id="runAtsBtn"' in index
    assert 'id="runAtsBtn" class="primary-button" type="button" disabled' in index
    assert "def select_ats_file" in bridge
    assert "def analyze_ats_file" in bridge
    assert re.search(r"analyze_ats_file\s*\(\s*cvmV101\.atsFile\.path\s*\)", v101)


def test_file_based_ats_checker_reads_docx() -> None:
    with tempfile.TemporaryDirectory(prefix="cvm_v101_ats_") as temp_dir:
        path = Path(temp_dir) / "candidate.docx"
        document = Document()
        document.add_heading("Alex Morgan", level=1)
        document.add_paragraph("alex.morgan@example.com | +33 6 12 34 56 78 | linkedin.com/in/alexmorgan")
        document.add_heading("Professional Summary", level=2)
        document.add_paragraph(
            "Data analyst focused on reporting, process improvement, stakeholder support, and reliable decision-making through clear data products."
        )
        document.add_heading("Professional Experience", level=2)
        document.add_paragraph("Data Analyst - Example Company")
        document.add_paragraph("- Automated recurring reporting and improved delivery time by 30%.")
        document.add_paragraph("- Built dashboards used by operational teams for weekly performance reviews.")
        document.add_heading("Education", level=2)
        document.add_paragraph("Master in Business Analytics")
        document.add_heading("Skills", level=2)
        document.add_paragraph("SQL, Python, Power BI, Excel, Tableau")
        document.save(path)

        result = run_ats_file_check(path)

        assert result["file_type"] == "DOCX"
        assert result["file_name"] == "candidate.docx"
        assert 0 <= result["score"] <= 100
        assert result["checks_run"] > 0


def test_editor_toolbar_and_native_save_as_patch() -> None:
    bridge = (ROOT / "src" / "bridge.py").read_text(encoding="utf-8")
    v101 = (ROOT / "ui" / "js" / "v101_core.js").read_text(encoding="utf-8")

    assert "def save_export_as" in bridge
    assert "webview.FileDialog.SAVE" in bridge
    assert "cvmV101CustomizeEditorToolbar" in v101
    assert '"#cvmV06CoachBtn"' in v101
    assert '"#cvmV06PdfBtn"' in v101
    assert '"#cvmV06DocxBtn"' in v101
    assert "cvmV101OpenSaveAsModal" in v101


def test_support_resume_and_log_copy_behaviour_are_present() -> None:
    index = (ROOT / "ui" / "index.html").read_text(encoding="utf-8")
    v101 = (ROOT / "ui" / "js" / "v101_core.js").read_text(encoding="utf-8")

    assert 'id="cvmResumeEditorBtn"' in index
    assert "cvmV101FixSupportButton" in v101
    assert re.search(r'cvmShowModal\s*\(\s*["\']cvmSupportModal["\']\s*\)', v101)
    assert "cvmV101EnsureResumeButton" in v101
    assert "cvmV101EnhanceLogCopyButtons" in v101

