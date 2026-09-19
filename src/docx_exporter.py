from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.design_utils import section_enabled, target_color, valid_photo_path
from src.models import CVProject, normalize_section_order
from src.i18n import normalize_language, tr
from src.project_manager import safe_filename
from src.cv_engine import format_date_range, format_month_year, parse_month_year


def rgb(value: str) -> RGBColor:
    value = value.replace("#", "")
    return RGBColor(
        int(value[0:2], 16),
        int(value[2:4], 16),
        int(value[4:6], 16),
    )


def add_bottom_border(
    paragraph,
    color: str,
) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    paragraph_borders = paragraph_properties.find(qn("w:pBdr"))

    if paragraph_borders is None:
        paragraph_borders = OxmlElement("w:pBdr")
        paragraph_properties.append(paragraph_borders)

    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "single")
    bottom_border.set(qn("w:sz"), "8")
    bottom_border.set(qn("w:space"), "2")
    bottom_border.set(qn("w:color"), color.replace("#", ""))
    paragraph_borders.append(bottom_border)



class DOCXExporter:
    def _label(self, project: CVProject, key: str) -> str:
        language = normalize_language(
            getattr(project.design, "cv_language", "") or project.design.language
        )
        section_lookup = {
            "profile": "profile",
            "experience": "experiences",
            "education": "education",
            "projects": "projects",
            "volunteering": "volunteering",
            "skills": "skills",
            "languages": "languages",
            "certifications": "certifications",
        }
        section_key = section_lookup.get(key)
        if section_key:
            custom = str(
                getattr(project.design, "section_titles", {}).get(section_key, "")
            ).strip()
            if custom:
                return custom
        return tr(language, key)

    def export(
        self,
        project: CVProject,
        output_target: Path,
    ) -> dict:
        output_target = Path(output_target)
        document = Document()

        self._active_project = project
        self._active_section = ""
        self._configure_document(document, project)
        self._render_header(document, project)

        renderers = {
            "profile": self._render_profile,
            "experiences": self._render_experiences,
            "education": self._render_education,
            "projects": self._render_projects,
            "volunteering": self._render_volunteering,
            "skills": self._render_skills,
            "languages": self._render_languages,
            "certifications": self._render_certifications,
        }

        for section_key in normalize_section_order(
            project.design.section_order
        ):
            if not section_enabled(project.design, section_key):
                continue

            renderer = renderers.get(section_key)

            if renderer:
                renderer(document, project)

        self._apply_default_text_color(document)

        base_name = safe_filename(
            project.personal.full_name
            or project.title
            or "CV"
        )

        if output_target.suffix.lower() == ".docx":
            output_path = output_target
        else:
            output_path = output_target / f"{base_name}_CV.docx"

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(output_path)

        return {
            "ok": True,
            "docx_path": str(output_path),
        }

    def _configure_document(self, document: Document, project: CVProject) -> None:
        section = document.sections[0]
        spacing = project.design.spacing

        section.top_margin = Cm(spacing["page_top_mm"] / 10)
        section.bottom_margin = Cm(spacing["page_bottom_mm"] / 10)
        section.left_margin = Cm(spacing["page_left_mm"] / 10)
        section.right_margin = Cm(spacing["page_right_mm"] / 10)

        normal_style = document.styles["Normal"]
        normal_style.font.name = "Arial"
        normal_style.font.size = Pt(8.2)
        normal_style.paragraph_format.space_after = Pt(0)
        normal_style.paragraph_format.space_before = Pt(0)
        normal_style.paragraph_format.line_spacing = spacing["body_line_height"]

    @staticmethod
    def _mm_to_pt(value: float) -> float:
        return float(value) * 2.834645669

    def _spacing(self, key: str, default: float = 0.0) -> float:
        project = getattr(self, "_active_project", None)

        if not project:
            return default

        try:
            return float(project.design.spacing.get(key, default))
        except (TypeError, ValueError, AttributeError):
            return default

    def _apply_section_line_spacing(self, paragraph) -> None:
        section = getattr(self, "_active_section", "")

        if not section:
            return

        line_height = self._spacing(
            f"{section}_line_height",
            self._spacing("body_line_height", 1.11),
        )
        paragraph.paragraph_format.line_spacing = line_height

    def _color_run(
        self,
        run,
        project: CVProject,
        target: str,
    ) -> None:
        run.font.color.rgb = rgb(
            target_color(
                project.design,
                target,
            )
        )

    def _add_section_title(
        self,
        document: Document,
        project: CVProject,
        text: str,
    ) -> None:
        paragraph = document.add_paragraph()

        section_gap = self._spacing("section_gap_mm", 1.0)
        title_before = self._spacing("title_before_mm", 0.0)
        title_after = self._spacing("title_after_mm", 0.6)

        paragraph.paragraph_format.space_before = Pt(
            self._mm_to_pt(section_gap + title_before)
        )
        paragraph.paragraph_format.space_after = Pt(
            self._mm_to_pt(title_after)
        )

        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        self._color_run(run, project, "section_titles")

        add_bottom_border(
            paragraph,
            target_color(
                project.design,
                "lines",
            ),
        )

    def _add_bullet(
        self,
        document: Document,
        text: str,
    ) -> None:
        text = (text or "").strip()

        if not text:
            return

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.28)
        paragraph.paragraph_format.first_line_indent = Cm(-0.22)
        paragraph.paragraph_format.space_after = Pt(
            self._mm_to_pt(
                self._spacing("bullet_gap_mm", 0.2)
            )
        )
        self._apply_section_line_spacing(paragraph)

        run = paragraph.add_run(f"• {text}")
        run.font.size = Pt(8)

    def _add_tools(
        self,
        document: Document,
        project: CVProject,
        tools: list[str],
    ) -> None:
        if not tools:
            return

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)

        label = paragraph.add_run(f"{self._label(project, 'tools')}: ")
        label.bold = True
        label.font.size = Pt(7.9)

        values = paragraph.add_run(", ".join(tools))
        values.font.size = Pt(7.9)

    def _render_header(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        personal = project.personal
        photo = valid_photo_path(project.design)

        if photo:
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = Cm(15.5)
            table.columns[1].width = Cm(2.5)

            left = table.cell(0, 0)
            right = table.cell(0, 1)
            left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

            left_paragraph = left.paragraphs[0]
            left_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._render_header_text(left_paragraph, project)

            photo_paragraph = right.paragraphs[0]
            photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = photo_paragraph.add_run()

            try:
                run.add_picture(
                    str(photo),
                    width=Cm(2.25),
                )
            except Exception:
                pass

            # Remove table borders. The table exists only to visually position
            # the optional photo and is flagged by the ATS checker as a risk.
            tbl = table._tbl
            tblPr = tbl.tblPr
            borders = OxmlElement("w:tblBorders")

            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = OxmlElement(f"w:{edge}")
                tag.set(qn("w:val"), "nil")
                borders.append(tag)

            tblPr.append(borders)

        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._render_header_text(paragraph, project)

        line = document.add_paragraph()
        add_bottom_border(
            line,
            target_color(
                project.design,
                "lines",
            ),
        )

    def _render_header_text(
        self,
        paragraph,
        project: CVProject,
    ) -> None:
        personal = project.personal

        name = paragraph.add_run(
            personal.full_name or "YOUR NAME"
        )
        name.bold = True
        name.font.size = Pt(20)
        self._color_run(name, project, "name")
        name.add_break()

        roles = [
            value
            for value in (
                personal.target_role,
                personal.secondary_target_role,
            )
            if value
        ]

        if roles:
            role = paragraph.add_run(" | ".join(roles))
            role.bold = True
            role.font.size = Pt(9.5)
            self._color_run(role, project, "role_titles")
            role.add_break()

        contacts = [
            personal.location,
            personal.email,
            personal.phone,
            "LinkedIn" if personal.linkedin else "",
            "GitHub" if personal.github else "",
            personal.website,
        ]

        contacts = [item for item in contacts if item]

        if contacts:
            contact = paragraph.add_run(" | ".join(contacts))
            contact.font.size = Pt(7.5)
            self._color_run(contact, project, "links")

    def _render_profile(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "profile"
        if not project.profile.strip():
            return

        self._add_section_title(document, project, self._label(project, "profile"))
        paragraph = document.add_paragraph(project.profile)
        paragraph.paragraph_format.space_after = Pt(
            self._mm_to_pt(
                self._spacing("profile_after_mm", 0.0)
            )
        )
        self._apply_section_line_spacing(paragraph)

    def _render_experiences(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "experiences"
        if not project.experiences:
            return

        self._add_section_title(
            document,
            project,
            self._label(project, "experience"),
        )

        for item in project.experiences:
            title = document.add_paragraph()
            self._apply_section_line_spacing(title)
            run = title.add_run(item.job_title)
            run.bold = True
            run.font.size = Pt(9)
            self._color_run(run, project, "role_titles")

            info = document.add_paragraph()
            self._apply_section_line_spacing(info)
            company = info.add_run(item.company)
            company.bold = True
            company.font.size = Pt(8.3)
            self._color_run(company, project, "subtitles")

            language = getattr(project.design, "cv_language", "") or project.design.language
            dates = format_date_range(
                item.start_date,
                item.end_date,
                current=item.current,
                language=language,
                present_label=self._label(project, "present"),
            )

            if dates:
                date_run = info.add_run(f" | {dates}")
                date_run.bold = True
                date_run.font.size = Pt(8.3)
                self._color_run(date_run, project, "subtitles")

            if item.location:
                location = document.add_paragraph(item.location)
                self._apply_section_line_spacing(location)
                if location.runs:
                    location.runs[0].font.size = Pt(7.7)

            if (
                item.description_mode == "paragraph"
                and item.paragraph.strip()
            ):
                paragraph = document.add_paragraph(item.paragraph)
                self._apply_section_line_spacing(paragraph)
                if paragraph.runs:
                    paragraph.runs[0].font.size = Pt(8)
            else:
                for bullet in item.bullets:
                    self._add_bullet(document, bullet)

            self._add_tools(document, project, item.tools)

    def _render_education(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "education"
        if not project.education:
            return

        self._add_section_title(document, project, self._label(project, "education"))

        for item in project.education:
            title = document.add_paragraph()
            self._apply_section_line_spacing(title)
            run = title.add_run(item.degree)
            run.bold = True
            run.font.size = Pt(8.5)
            self._color_run(run, project, "role_titles")

            details = ", ".join(
                value
                for value in (
                    item.school,
                    item.location,
                )
                if value
            )

            language = getattr(project.design, "cv_language", "") or project.design.language
            dates = format_date_range(
                item.start_date,
                item.end_date,
                language=language,
                present_label=self._label(project, "present"),
            )

            if dates:
                details = f"{details} | {dates}" if details else dates

            if details:
                paragraph = document.add_paragraph()
                self._apply_section_line_spacing(paragraph)
                run = paragraph.add_run(details)
                run.font.size = Pt(7.8)
                self._color_run(run, project, "subtitles")

            if item.details:
                paragraph = document.add_paragraph(item.details)
                self._apply_section_line_spacing(paragraph)
                if paragraph.runs:
                    paragraph.runs[0].font.size = Pt(7.7)

    def _render_projects(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "projects"
        if not project.projects:
            return

        self._add_section_title(document, project, self._label(project, "projects"))

        for item in project.projects:
            title = document.add_paragraph()
            self._apply_section_line_spacing(title)
            run = title.add_run(item.name)
            run.bold = True
            run.font.size = Pt(8.5)
            self._color_run(run, project, "role_titles")

            language = getattr(project.design, "cv_language", "") or project.design.language

            # CVM 0.6.1: projects show one year only.
            dates = str(getattr(item, "year", "") or "").strip()

            if not dates:
                # Backward compatibility for older project files.
                for legacy_value in (
                    getattr(item, "date", ""),
                    getattr(item, "end_date", ""),
                    getattr(item, "start_date", ""),
                ):
                    legacy_year, _legacy_month = parse_month_year(
                        legacy_value
                    )

                    if legacy_year is not None:
                        dates = str(legacy_year)
                        break

            if dates:
                date_run = title.add_run(f" | {dates}")
                date_run.bold = True
                date_run.font.size = Pt(8.5)

            if item.subtitle:
                paragraph = document.add_paragraph()
                self._apply_section_line_spacing(paragraph)
                run = paragraph.add_run(item.subtitle)
                run.font.size = Pt(7.8)
                self._color_run(run, project, "subtitles")

            for bullet in item.bullets:
                self._add_bullet(document, bullet)

            self._add_tools(document, project, item.tools)

            if item.link:
                paragraph = document.add_paragraph()
                self._apply_section_line_spacing(paragraph)
                run = paragraph.add_run(item.link)
                run.font.size = Pt(7.5)
                self._color_run(run, project, "links")

    def _render_volunteering(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "volunteering"
        if not project.volunteering:
            return

        self._add_section_title(document, project, self._label(project, "volunteering"))

        for item in project.volunteering:
            title = document.add_paragraph()
            self._apply_section_line_spacing(title)
            run = title.add_run(item.role)
            run.bold = True
            run.font.size = Pt(8.5)
            self._color_run(run, project, "role_titles")

            details = ", ".join(
                value
                for value in (
                    item.organization,
                    item.location,
                )
                if value
            )

            language = getattr(project.design, "cv_language", "") or project.design.language
            dates = format_date_range(
                item.start_date,
                item.end_date,
                language=language,
                present_label=self._label(project, "present"),
            )

            if dates:
                details = f"{details} | {dates}" if details else dates

            if details:
                paragraph = document.add_paragraph()
                self._apply_section_line_spacing(paragraph)
                run = paragraph.add_run(details)
                run.font.size = Pt(7.8)
                self._color_run(run, project, "subtitles")

            for bullet in item.bullets:
                self._add_bullet(document, bullet)

    def _render_skills(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "skills"
        groups = [
            group
            for group in project.skills
            if group.name or group.items
        ]

        if not groups:
            return

        self._add_section_title(document, project, self._label(project, "skills"))

        for group in groups:
            paragraph = document.add_paragraph()

            if group.name:
                label = paragraph.add_run(f"{group.name}: ")
                label.bold = True
                label.font.size = Pt(8)

            values = paragraph.add_run(", ".join(group.items))
            values.font.size = Pt(8)

    def _render_languages(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "languages"
        if not project.languages:
            return

        self._add_section_title(document, project, self._label(project, "languages"))
        entries = []

        for item in project.languages:
            if item.name and item.level:
                entries.append(f"{item.name}: {item.level}")
            elif item.name:
                entries.append(item.name)

        if entries:
            paragraph = document.add_paragraph(" | ".join(entries))
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(8)

    def _render_certifications(
        self,
        document: Document,
        project: CVProject,
    ) -> None:
        self._active_section = "certifications"
        if not project.certifications:
            return

        self._add_section_title(document, project, self._label(project, "certifications"))
        entries = []

        for item in project.certifications:
            text = item.name

            if item.issuer:
                text = f"{text} - {item.issuer}" if text else item.issuer

            if item.year:
                text += f" ({item.year})"

            if text:
                entries.append(text)

        if entries:
            paragraph = document.add_paragraph(" | ".join(entries))
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(7.9)

    def _apply_default_text_color(
        self,
        document: Document,
    ) -> None:
        default_color = RGBColor(17, 24, 39)

        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                if run.font.color.rgb is None:
                    run.font.color.rgb = default_color
