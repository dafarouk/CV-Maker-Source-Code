from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Any

from src.design_utils import section_enabled, valid_photo_path
from src.models import CVProject


ACTION_VERBS = {
    "analysed", "analyzed", "automated", "built", "created",
    "developed", "designed", "implemented", "improved", "managed",
    "monitored", "optimized", "produced", "reduced", "reported",
    "streamlined", "validated", "coordinated", "delivered", "identified",
    "measured", "transformed", "maintained", "supported", "led",
}

STANDARD_SECTIONS = {
    "profile",
    "experiences",
    "education",
    "projects",
    "volunteering",
    "skills",
    "languages",
    "certifications",
}

SUSPICIOUS_GLYPHS = {
    "★", "☆", "◆", "◇", "●", "○", "►", "▶", "✓", "✔", "✦", "✧",
}


def _words(text: str) -> list[str]:
    return re.findall(
        r"[A-Za-zÀ-ÿ0-9+#.\-]+",
        (text or "").lower(),
    )


def _emails(text: str) -> list[str]:
    return re.findall(
        r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
        text or "",
    )


def _all_cv_text(project: CVProject) -> str:
    chunks = [
        project.personal.full_name,
        project.personal.target_role,
        project.personal.secondary_target_role,
        project.personal.location,
        project.personal.email,
        project.personal.phone,
        project.profile,
    ]

    for item in project.experiences:
        chunks.extend([
            item.job_title,
            item.company,
            item.location,
            item.paragraph,
            *item.bullets,
            *item.tools,
        ])

    for item in project.education:
        chunks.extend([
            item.degree,
            item.school,
            item.location,
            item.details,
        ])

    for item in project.projects:
        chunks.extend([
            item.name,
            item.subtitle,
            *item.bullets,
            *item.tools,
        ])

    for item in project.volunteering:
        chunks.extend([
            item.role,
            item.organization,
            *item.bullets,
        ])

    for group in project.skills:
        chunks.extend([group.name, *group.items])

    for item in project.languages:
        chunks.extend([item.name, item.level])

    for item in project.certifications:
        chunks.extend([item.name, item.issuer, item.year])

    return "\n".join(str(value) for value in chunks if value)


def run_ats_check(
    project: CVProject,
    pdf_audit: dict[str, Any] | None = None,
) -> dict:
    deductions = 0
    issues: list[str] = []
    warnings: list[str] = []
    passed: list[str] = []
    technical: list[str] = []

    personal = project.personal
    full_text = _all_cv_text(project)

    # --------------------------------------------------------
    # CONTACT / IDENTITY
    # --------------------------------------------------------

    if personal.full_name.strip():
        passed.append("Candidate name is present.")
    else:
        deductions += 14
        issues.append("Add your full name.")

    if personal.email.strip():
        if _emails(personal.email):
            passed.append("Email address has a recognisable format.")
        else:
            deductions += 5
            issues.append("The email address format looks invalid.")
    else:
        deductions += 9
        issues.append("Add an email address.")

    if personal.phone.strip():
        digits = re.sub(r"\D", "", personal.phone)

        if len(digits) >= 8:
            passed.append("Phone number is present.")
        else:
            deductions += 3
            warnings.append("Phone number looks unusually short.")
    else:
        deductions += 3
        warnings.append("Phone number is missing.")

    if personal.target_role.strip():
        passed.append("A target role is clearly stated.")
    else:
        deductions += 4
        warnings.append("Add a clear target role near the top of the CV.")

    # --------------------------------------------------------
    # CONTENT STRUCTURE
    # --------------------------------------------------------

    profile_words = len(_words(project.profile))

    if project.profile and section_enabled(project.design, "profile"):
        if 18 <= profile_words <= 90:
            passed.append("Profile length is concise and readable.")
        elif profile_words > 90:
            deductions += 3
            warnings.append("Profile is long; aim for roughly 2–5 concise lines.")
        else:
            deductions += 2
            warnings.append("Profile is very short and may not communicate enough context.")

    if project.experiences and section_enabled(project.design, "experiences"):
        passed.append("Professional experience section is present.")
    else:
        deductions += 10
        warnings.append("No visible professional experience section is present.")

    if project.education and section_enabled(project.design, "education"):
        passed.append("Education section is present.")
    else:
        deductions += 4
        warnings.append("Education section is empty or hidden.")

    skill_count = sum(len(group.items) for group in project.skills)

    if skill_count >= 5 and section_enabled(project.design, "skills"):
        passed.append("Technical skills are explicitly listed.")
    elif skill_count:
        deductions += 3
        warnings.append("The visible skills section is relatively sparse.")
    else:
        deductions += 7
        issues.append("Technical skills section is empty.")

    # --------------------------------------------------------
    # BULLETS / IMPACT / REPETITION
    # --------------------------------------------------------

    all_bullets: list[str] = []
    long_bullets = 0
    very_short_bullets = 0
    action_verb_misses = 0
    quantified_bullets = 0

    for experience in project.experiences:
        if not experience.job_title.strip():
            deductions += 2
            issues.append("An experience is missing its job title.")

        if not experience.company.strip():
            deductions += 2
            issues.append("An experience is missing its company name.")

        if not experience.start_date.strip():
            deductions += 1
            warnings.append("An experience is missing its start date.")

        for bullet in experience.bullets:
            bullet = bullet.strip()

            if not bullet:
                continue

            all_bullets.append(bullet)
            word_count = len(_words(bullet))

            if word_count > 38:
                long_bullets += 1

            if 0 < word_count < 6:
                very_short_bullets += 1

            bullet_words = _words(bullet)

            if bullet_words and bullet_words[0] not in ACTION_VERBS:
                action_verb_misses += 1

            if re.search(r"\b\d+(?:[.,]\d+)?\s*(?:%|€|\$|£|k|m|hours?|days?|users?|reports?|dashboards?|records?|files?)?\b", bullet, re.I):
                quantified_bullets += 1

    if long_bullets:
        deductions += min(6, long_bullets)
        warnings.append(f"{long_bullets} experience bullet(s) are long and may scan poorly.")

    if very_short_bullets:
        deductions += min(4, very_short_bullets)
        warnings.append(f"{very_short_bullets} experience bullet(s) are extremely short.")

    if action_verb_misses and all_bullets:
        warnings.append(
            f"{action_verb_misses} experience bullet(s) do not begin with one of CVM's recognised action verbs."
        )
    elif all_bullets:
        passed.append("Experience bullets generally start with clear action verbs.")

    if quantified_bullets:
        passed.append(f"{quantified_bullets} experience bullet(s) contain measurable evidence.")
    elif all_bullets:
        warnings.append("No measurable results were detected in experience bullets; add numbers only where they are real.")

    normalized_bullets = [
        re.sub(r"\W+", " ", bullet.lower()).strip()
        for bullet in all_bullets
    ]

    duplicates = len(normalized_bullets) - len(set(normalized_bullets))

    if duplicates:
        deductions += min(6, duplicates * 2)
        warnings.append(f"{duplicates} duplicated experience bullet(s) detected.")
    elif normalized_bullets:
        passed.append("No identical experience bullets were detected.")

    # Repeated leading verbs can make a CV read mechanically.
    first_words = [
        words[0]
        for bullet in all_bullets
        if (words := _words(bullet))
    ]

    repeated_starts = [
        word
        for word, count in Counter(first_words).items()
        if count >= 3
    ]

    if repeated_starts:
        warnings.append(
            "Several bullets repeatedly start with: "
            + ", ".join(sorted(repeated_starts))
            + ". Consider varying wording when truthful."
        )

    # --------------------------------------------------------
    # ATS-SAFETY / FORMAT CHOICES
    # --------------------------------------------------------

    if project.design.template_id == "ats_classic":
        passed.append("ATS Classic uses a single-column text-first structure.")

    if project.design.photo_enabled:
        if valid_photo_path(project.design):
            deductions += 2
            warnings.append(
                "A photo is enabled. CVM keeps it outside the main text flow, but some employers and regions prefer photo-free CVs."
            )
        else:
            deductions += 3
            issues.append("Photo is enabled but the selected photo file is unavailable or unsupported.")

    suspicious = sorted(
        glyph
        for glyph in SUSPICIOUS_GLYPHS
        if glyph in full_text
    )

    if suspicious:
        deductions += min(4, len(suspicious))
        warnings.append(
            "Decorative symbols detected in CV text: "
            + " ".join(suspicious)
            + ". Plain text is usually safer for parsing."
        )
    else:
        passed.append("No risky decorative glyphs were detected in the CV text.")

    visible_sections = [
        key
        for key, value in project.design.section_visibility.items()
        if value
    ]

    unknown_sections = [
        key
        for key in visible_sections
        if key not in STANDARD_SECTIONS
    ]

    if unknown_sections:
        deductions += 2
        warnings.append("Non-standard section identifiers were detected.")
    else:
        passed.append("Section structure uses conventional CV categories.")

    # --------------------------------------------------------
    # EXPORTED PDF PARSING AUDIT
    # --------------------------------------------------------

    if pdf_audit:
        if not pdf_audit.get("available", True):
            warnings.append(pdf_audit.get("message", "PDF parsing audit was unavailable."))
        elif not pdf_audit.get("compiled", True):
            deductions += 8
            issues.append(pdf_audit.get("message", "The CV could not be compiled for parsing verification."))
        else:
            if pdf_audit.get("selectable_text"):
                passed.append("Generated PDF contains selectable/extractable text.")
            else:
                deductions += 18
                issues.append("Generated PDF text could not be reliably extracted.")

            if pdf_audit.get("name_found"):
                passed.append("Candidate name is recoverable from the generated PDF text layer.")
            else:
                deductions += 10
                issues.append("Candidate name was not found in extracted PDF text.")

            expected = pdf_audit.get("expected_sections", [])
            found = pdf_audit.get("found_sections", [])

            if expected:
                coverage = len(found) / len(expected)

                if coverage >= 0.95:
                    passed.append("Visible CV section headings were recovered from the PDF text layer.")
                elif coverage >= 0.70:
                    deductions += 5
                    warnings.append("Some expected section headings were not recovered from the PDF text layer.")
                else:
                    deductions += 10
                    issues.append("Several expected section headings were missing from extracted PDF text.")

            if pdf_audit.get("section_order_valid"):
                passed.append("Extracted PDF section order is consistent with the intended reading order.")
            else:
                deductions += 8
                issues.append("Extracted PDF section order may not match the intended reading order.")

            page_count = int(pdf_audit.get("page_count", 0) or 0)
            desired_pages = 1 if project.design.page_mode == "1" else 2

            if page_count and page_count <= desired_pages:
                passed.append(f"Generated PDF fits the selected {desired_pages}-page target.")
            elif page_count:
                deductions += 4
                warnings.append(
                    f"Generated PDF currently uses {page_count} page(s), above the selected {desired_pages}-page target."
                )

            technical.append(
                f"PDF extraction: {pdf_audit.get('characters_extracted', 0)} characters, "
                f"{pdf_audit.get('block_count', 0)} text blocks, "
                f"{page_count} page(s)."
            )

    score = max(0, min(100, 100 - deductions))

    if score >= 90:
        label = "Strong"
    elif score >= 78:
        label = "Good"
    elif score >= 65:
        label = "Needs review"
    else:
        label = "Needs work"

    return {
        "score": score,
        "label": label,
        "issues": issues,
        "warnings": warnings,
        "passed": passed,
        "technical": technical,
        "checks_run": len(issues) + len(warnings) + len(passed),
        "note": (
            "CVM performs a local ATS-readiness audit of content, structure and the generated PDF text layer. "
            "It is not a score from Workday, Taleo, Greenhouse, SuccessFactors or any other employer ATS."
        ),
    }


# ============================================================
# FILE-BASED ATS CHECK (PDF / DOCX)
# ============================================================

ATS_FILE_EXTENSIONS = {".pdf", ".docx"}

_SECTION_ALIASES = {
    "profile": (
        "profile", "professional profile", "summary", "professional summary",
        "profil", "profil professionnel", "résumé", "resume",
    ),
    "experience": (
        "experience", "professional experience", "work experience",
        "employment", "expérience", "expériences", "expérience professionnelle",
        "expériences professionnelles",
    ),
    "education": (
        "education", "academic background", "studies", "formation",
        "formations", "parcours académique",
    ),
    "skills": (
        "skills", "technical skills", "core skills", "competencies",
        "compétences", "compétences techniques",
    ),
}


def _normalize_heading_line(value: str) -> str:
    value = re.sub(r"[^A-Za-zÀ-ÿ ]+", " ", value or "")
    return re.sub(r"\s+", " ", value).strip().lower()


def _contains_section(text: str, aliases: tuple[str, ...]) -> bool:
    lines = [_normalize_heading_line(line) for line in (text or "").splitlines()]

    for line in lines:
        if not line or len(line) > 70:
            continue

        if any(line == alias or line.startswith(alias + " ") for alias in aliases):
            return True

    return False


def _extract_pdf_text(path: Path) -> dict[str, Any]:
    import fitz

    pages_text: list[str] = []
    block_count = 0

    with fitz.open(path) as document:
        page_count = document.page_count

        for page in document:
            page_text = page.get_text("text") or ""
            pages_text.append(page_text)
            block_count += len(page.get_text("blocks") or [])

    text = "\n".join(pages_text).strip()

    return {
        "text": text,
        "page_count": page_count,
        "block_count": block_count,
        "table_count": 0,
        "format": "PDF",
        "selectable_text": bool(text and len(_words(text)) >= 20),
    }


def _extract_docx_text(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(str(path))
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            chunks.append(value)

    table_count = len(document.tables)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                chunks.append(" | ".join(values))

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            value = paragraph.text.strip()
            if value:
                chunks.append(value)

        for paragraph in section.footer.paragraphs:
            value = paragraph.text.strip()
            if value:
                chunks.append(value)

    text = "\n".join(chunks).strip()

    return {
        "text": text,
        "page_count": 0,
        "block_count": len(document.paragraphs),
        "table_count": table_count,
        "format": "DOCX",
        "selectable_text": bool(text and len(_words(text)) >= 20),
    }


def _extract_cv_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf_text(path)

    if suffix == ".docx":
        return _extract_docx_text(path)

    raise ValueError("ATS Check supports PDF and DOCX files only.")


def run_ats_file_check(file_path: str | Path) -> dict:
    """
    Analyse a user-selected PDF or DOCX CV.

    This deliberately operates on the document itself instead of CVM's current
    structured editor project. The result keeps the same UI contract as the
    original structured ATS checker: score, label, issues, warnings, passed,
    technical, checks_run and note.
    """

    path = Path(file_path).expanduser().resolve()

    if not path.exists() or not path.is_file():
        raise ValueError("The selected CV file does not exist.")

    if path.suffix.lower() not in ATS_FILE_EXTENSIONS:
        raise ValueError("ATS Check supports PDF and DOCX files only.")

    extracted = _extract_cv_file(path)
    text = extracted["text"]
    words = _words(text)
    word_count = len(words)

    deductions = 0
    issues: list[str] = []
    warnings: list[str] = []
    passed: list[str] = []
    technical: list[str] = []

    # --------------------------------------------------------
    # TEXT EXTRACTION / BASIC PARSABILITY
    # --------------------------------------------------------

    if extracted["selectable_text"]:
        passed.append(f"{extracted['format']} contains extractable text for ATS parsing.")
    else:
        deductions += 30
        issues.append(
            "Very little extractable text was found. The CV may be image-based, scanned, empty or difficult for an ATS to parse."
        )

    if word_count >= 150:
        passed.append("The document contains enough text for a meaningful CV structure check.")
    elif word_count >= 60:
        deductions += 6
        warnings.append("The CV contains relatively little extractable text.")
    else:
        deductions += 12
        issues.append("The CV contains too little extractable text for a reliable ATS-readiness review.")

    if word_count > 1300:
        deductions += 4
        warnings.append("The CV is unusually text-heavy; review density and relevance.")

    # --------------------------------------------------------
    # CONTACT INFORMATION
    # --------------------------------------------------------

    emails = _emails(text)

    if emails:
        passed.append("An email address is recoverable from the document text.")
    else:
        deductions += 10
        issues.append("No email address was detected in the extracted CV text.")

    phone_candidates = re.findall(r"(?:\+?\d[\d\s().-]{6,}\d)", text)
    phone_candidates = [
        candidate
        for candidate in phone_candidates
        if 8 <= len(re.sub(r"\D", "", candidate)) <= 16
    ]

    if phone_candidates:
        passed.append("A phone number is recoverable from the document text.")
    else:
        deductions += 4
        warnings.append("No clear phone number was detected in the extracted CV text.")

    if re.search(r"linkedin(?:\.com)?|https?://|www\.", text, re.I):
        passed.append("A professional web or LinkedIn reference is present in the document text.")

    # --------------------------------------------------------
    # SECTION STRUCTURE
    # --------------------------------------------------------

    found_sections: list[str] = []

    for section_name, aliases in _SECTION_ALIASES.items():
        if _contains_section(text, aliases):
            found_sections.append(section_name)
            passed.append(f"A recognisable {section_name} section was detected.")
        else:
            deductions += 4
            warnings.append(f"No clear {section_name} section heading was detected.")

    if len(found_sections) >= 3:
        passed.append("The CV uses several conventional section headings that are easy to identify.")

    # --------------------------------------------------------
    # BULLET / CONTENT READABILITY
    # --------------------------------------------------------

    bullet_lines = [
        line.strip()
        for line in text.splitlines()
        if re.match(r"^\s*(?:[-–—•▪◦]|\d+[.)])\s+", line)
    ]

    long_bullets = [
        line
        for line in bullet_lines
        if len(_words(line)) > 40
    ]

    if bullet_lines:
        passed.append(f"{len(bullet_lines)} bullet-style line(s) were recovered from the document.")

    if long_bullets:
        deductions += min(6, len(long_bullets))
        warnings.append(
            f"{len(long_bullets)} extracted bullet(s) are long and may be harder to scan quickly."
        )

    suspicious = sorted(glyph for glyph in SUSPICIOUS_GLYPHS if glyph in text)

    if suspicious:
        deductions += min(4, len(suspicious))
        warnings.append(
            "Decorative symbols were detected in the extracted text: "
            + " ".join(suspicious)
            + ". Plain text is generally safer for ATS parsing."
        )
    else:
        passed.append("No risky decorative glyphs were detected in the extracted text.")

    # --------------------------------------------------------
    # FORMAT-SPECIFIC CHECKS
    # --------------------------------------------------------

    if extracted["format"] == "PDF":
        page_count = int(extracted.get("page_count", 0) or 0)

        if 1 <= page_count <= 2:
            passed.append(f"PDF length is {page_count} page(s).")
        elif page_count > 2:
            deductions += 4
            warnings.append(f"The PDF is {page_count} pages long; verify that all content is necessary.")

        technical.append(
            f"PDF extraction: {len(text)} characters, {word_count} words, "
            f"{extracted.get('block_count', 0)} text blocks, {page_count} page(s)."
        )

    else:
        table_count = int(extracted.get("table_count", 0) or 0)

        if table_count:
            deductions += min(6, table_count * 2)
            warnings.append(
                f"The DOCX contains {table_count} table(s). Complex table-based layouts can be less reliable in some ATS parsers."
            )
        else:
            passed.append("No Word tables were detected in the DOCX structure.")

        technical.append(
            f"DOCX extraction: {len(text)} characters, {word_count} words, "
            f"{extracted.get('block_count', 0)} paragraph block(s), {table_count} table(s)."
        )

    score = max(0, min(100, 100 - deductions))

    if score >= 90:
        label = "Strong"
    elif score >= 78:
        label = "Good"
    elif score >= 65:
        label = "Needs review"
    else:
        label = "Needs work"

    return {
        "score": score,
        "label": label,
        "issues": issues,
        "warnings": warnings,
        "passed": passed,
        "technical": technical,
        "checks_run": len(issues) + len(warnings) + len(passed),
        "file_name": path.name,
        "file_type": extracted["format"],
        "note": (
            "CVM performs a local ATS-readiness review of the selected PDF or DOCX using extracted text and document structure. "
            "It is not a score from Workday, Taleo, Greenhouse, SuccessFactors or any other employer ATS."
        ),
    }
